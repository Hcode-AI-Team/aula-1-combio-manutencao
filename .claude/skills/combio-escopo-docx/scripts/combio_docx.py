#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
combio_docx.py — Engine de montagem dos documentos de escopo e relatorio de
entrega da ComBio Energia, no layout canonico (anatomia medida no arquivo de
referencia: ver references/layout_anatomia.md).

Uso como biblioteca:

    from combio_docx import CombioDoc
    doc = CombioDoc(header="ESCOPO DE MELHORIA", eyebrow="SOLICITACAO DE MELHORIA",
                    titulo="Escopo de Melhoria")
    doc.faixa("IDENTIFICACAO")
    doc.campos([["Titulo da melhoria:", ""], ["Categoria:", "Sistema"]])
    doc.salvar("saida.docx")

Ou por spec JSON, via build_doc.py.

Regras duras implementadas aqui (nao alterar sem atualizar a skill):
  - A4 11906x16838 dxa; margens topo 1500 / base 1300 / laterais 1300;
    header e footer 708. Largura util de tabela: 9300 dxa.
  - Cabecalho: logo 1143000x247650 EMU a esquerda, tab direito em 9026,
    tipo do documento Kodchasan Bold 13pt #116533 caixa alta, borda inferior
    single #399444 sz=10 space=8.
  - Rodape: borda superior single #399444 sz=8 space=6, Varela Round 8pt
    #4E5B50, "ComBio Energia" + tab + "Pagina X de Y" (campos PAGE/NUMPAGES).
  - Faixa de secao: tabela 1x1 de 9300 dxa, shading CLEAR fill 116533, texto
    Kodchasan Bold 11pt branco, letter-spacing 12, caixa alta, keepNext.
  - Zebra de bloco de campos: FFFFFF / F2F4F1. Zebra de tabela de dados: F9FAF8.
  - Bordas de tabela: single #C5CCC6 sz=4.
  - Marcadores de status coloridos: OK #116533, ATENCAO #BF882D, ERRO #AB3434.
"""

from __future__ import annotations

import copy
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor, Twips

# --------------------------------------------------------------------------
# Tokens — fonte unica: combio-platform/design/colors_and_type.css
# NUNCA inventar hex fora desta lista.
# --------------------------------------------------------------------------
GREEN_DARK = "116533"   # faixas, cabecalho de tabela, titulo, ancora de diagrama
GREEN_LIGHT = "399444"  # eyebrow, reguas, destino de diagrama
GREEN_EARTH = "1C361F"
YELLOW = "BF882D"       # atencao / callout / barra de ressalva
RED = "AB3434"          # so erro real
BLUE_GREY = "76858E"    # neutro executivo / intermediario de diagrama
INK = "0F1A11"
INK_80 = "2A3A2D"
INK_60 = "4E5B50"
INK_20 = "C5CCC6"
INK_05 = "F2F4F1"
INK_02 = "F9FAF8"
WHITE = "FFFFFF"

FONT_PRIMARY = "Kodchasan"
FONT_PRIMARY_FALLBACK = "Trebuchet MS"
FONT_SECONDARY = "Varela Round"
FONT_MONO = "Consolas"

PAGE_W = 11906
PAGE_H = 16838
MARGIN_TOP = 1500
MARGIN_BOTTOM = 1300
MARGIN_SIDE = 1300
HEADER_DIST = 708
FOOTER_DIST = 708
CONTENT_W = 9300          # largura util de tabela, em dxa
TAB_RIGHT = 9026          # tab direito de cabecalho/rodape

LOGO_W_EMU = 1143000      # ~3,0 cm
LOGO_H_EMU = 247650       # ~0,65 cm

# Versao da skill — fonte unica. Ao alterar a skill, suba aqui e registre
# no CHANGELOG.md. A versao e carimbada nas propriedades de todo .docx gerado.
SKILL_NOME = "combio-escopo-docx"
VERSAO = "1.1.1"

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.normpath(os.path.join(HERE, "..", "assets"))
LOGO_PATH = os.path.join(ASSETS, "logo", "combio-primary.png")

STATUS_COLORS = {"✓": GREEN_DARK, "⚠": YELLOW, "✗": RED}
# Fonte dos marcadores de status. Precisa ser uma fonte de SIMBOLO (nao emoji),
# senao o "⚠" vira emoji colorido — proibido pelo padrao ComBio.
FONT_SIMBOLO = "Segoe UI Symbol"


# --------------------------------------------------------------------------
# helpers de XML cru
# --------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


# Ordem canonica dos filhos (ECMA-376). Inserir fora de ordem faz o Word
# "reparar" o arquivo e o LibreOffice simplesmente ignorar o elemento —
# foi o que quebrou os tab stops de cabecalho/rodape na primeira versao.
_ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
            "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
            "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
            "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
            "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
            "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl",
            "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps",
            "strike", "dstrike", "outline", "shadow", "emboss", "imprint",
            "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
            "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
            "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
             "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign",
             "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual",
              "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
              "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout",
              "tblCellMar", "tblLook"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore",
             "wAfter", "cantSplit", "trHeight", "tblHeader", "tblCellSpacing",
             "jc", "hidden"],
}


def _local(tag):
    return tag.split("}")[-1]


def _insert_ordered(parent, child):
    """Insere child em parent respeitando a ordem canonica do schema."""
    order = _ORDER.get(_local(parent.tag))
    if not order:
        parent.append(child)
        return child
    name = _local(child.tag)
    # remove duplicata existente
    for existing in parent.findall(qn("w:" + name)):
        parent.remove(existing)
    try:
        idx = order.index(name)
    except ValueError:
        parent.append(child)
        return child
    for sibling in parent:
        sname = _local(sibling.tag)
        if sname not in order or order.index(sname) > idx:
            sibling.addprevious(child)
            return child
    parent.append(child)
    return child


def _set_shading(element, fill):
    """Shading CLEAR (nunca SOLID) em tcPr ou pPr."""
    _insert_ordered(element, _el("w:shd", val="clear", color="auto", fill=fill))


def _borders(pr, color=INK_20, sz=4, sides=("top", "left", "bottom", "right"),
             val="single"):
    b = OxmlElement("w:tcBorders" if _local(pr.tag) == "tcPr" else "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if side in sides:
            b.append(_el("w:" + side, val=val, sz=sz, space=0, color=color))
        elif _local(pr.tag) == "tblPr":
            b.append(_el("w:" + side, val="nil"))
    _insert_ordered(pr, b)


def _cell_margins(tcPr, top=90, bottom=90, left=160, right=160):
    m = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        m.append(_el("w:" + side, w=val, type="dxa"))
    _insert_ordered(tcPr, m)


def _valign(tcPr, val="center"):
    _insert_ordered(tcPr, _el("w:vAlign", val=val))


def _keep_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def _field(paragraph, instr, font=FONT_SECONDARY, size=8, color=INK_60):
    """Insere um campo Word (PAGE, NUMPAGES)."""
    r1 = paragraph.add_run()
    r1._r.append(_el("w:fldChar", fldCharType="begin"))
    r2 = paragraph.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = " %s " % instr
    r2._r.append(t)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", fldCharType="end"))
    for r in (r1, r2, r3):
        _style_run(r, font=font, size=size, color=color)


def _style_run(run, font=FONT_PRIMARY, size=10.5, color=INK_80, bold=False,
               italic=False, caps=False, spacing=None, fallback=True):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        _insert_ordered(rPr, rFonts)
    for a in ("ascii", "hAnsi", "cs"):
        rFonts.set(qn("w:" + a), font)
    if caps:
        _insert_ordered(rPr, _el("w:caps", val="1"))
    if spacing is not None:
        _insert_ordered(rPr, _el("w:spacing", val=spacing))
    return run


def _p_spacing(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    if line:
        pf.line_spacing = line


def _para_border(paragraph, side="left", color=YELLOW, sz=24, space=8):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        _insert_ordered(pPr, pbdr)
    pbdr.append(_el("w:" + side, val="single", sz=sz, space=space, color=color))


def _para_shading(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    _set_shading(pPr, fill)


def _split_status(text):
    """Separa marcador de status inicial (checked/warn/cross) do resto."""
    if not text:
        return None, text
    ch = text[0]
    if ch in STATUS_COLORS:
        return ch, text[1:].lstrip()
    return None, text


# --------------------------------------------------------------------------
# numbering para bullets (nunca usar "-" ou bullet literal no texto)
# --------------------------------------------------------------------------
BULLET_ABSTRACT_XML = """
<w:abstractNum xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:abstractNumId="{aid}">
  <w:multiLevelType w:val="hybridMultilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>
    <w:lvlText w:val="●"/><w:lvlJc w:val="left"/>
    <w:pPr><w:ind w:left="360" w:hanging="200"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Kodchasan" w:hAnsi="Kodchasan" w:hint="default"/><w:sz w:val="16"/></w:rPr>
  </w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/>
    <w:lvlText w:val="○"/><w:lvlJc w:val="left"/>
    <w:pPr><w:ind w:left="720" w:hanging="200"/></w:pPr>
    <w:rPr><w:rFonts w:ascii="Kodchasan" w:hAnsi="Kodchasan" w:hint="default"/><w:sz w:val="16"/></w:rPr>
  </w:lvl>
</w:abstractNum>
"""


class CombioDoc:
    """Documento no layout canonico ComBio."""

    def __init__(self, header="ESCOPO", eyebrow=None, titulo=None,
                 subtitulo=None, logo=LOGO_PATH):
        self.doc = Document()
        self.logo = logo if (logo and os.path.exists(logo)) else None
        self._num_id = None
        self._setup_styles()
        self._setup_page()
        self._setup_header(header)
        self._setup_footer()
        if eyebrow:
            self.eyebrow(eyebrow)
        if titulo:
            self.titulo(titulo)
        if subtitulo:
            self.orientacao(subtitulo)

    # ---------------------------------------------------------------- setup
    def _setup_styles(self):
        st = self.doc.styles["Normal"]
        st.font.name = FONT_PRIMARY
        st.font.size = Pt(10.5)
        st.font.color.rgb = RGBColor.from_string(INK_80)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            _insert_ordered(rpr, rf)
        for a in ("ascii", "hAnsi", "cs"):
            rf.set(qn("w:" + a), FONT_PRIMARY)
        st.paragraph_format.space_after = Twips(90)
        st.paragraph_format.line_spacing = 1.15

        # Os estilos Header/Footer do template padrao trazem um tab CENTER
        # herdado; sem remove-lo, o "\t" para no centro em vez do tab direito.
        for nome in ("Header", "Footer"):
            try:
                el = self.doc.styles[nome].element
            except KeyError:
                continue
            pPr = el.find(qn("w:pPr"))
            if pPr is None:
                continue
            for tabs in pPr.findall(qn("w:tabs")):
                pPr.remove(tabs)

    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width = Twips(PAGE_W)
        s.page_height = Twips(PAGE_H)
        s.top_margin = Twips(MARGIN_TOP)
        s.bottom_margin = Twips(MARGIN_BOTTOM)
        s.left_margin = Twips(MARGIN_SIDE)
        s.right_margin = Twips(MARGIN_SIDE)
        s.header_distance = Twips(HEADER_DIST)
        s.footer_distance = Twips(FOOTER_DIST)

    def _setup_header(self, texto):
        hdr = self.doc.sections[0].header
        p = hdr.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Twips(TAB_RIGHT), WD_TAB_ALIGNMENT.RIGHT)
        _p_spacing(p, after=60)
        if self.logo:
            run = p.add_run()
            run.add_picture(self.logo, width=Emu(LOGO_W_EMU), height=Emu(LOGO_H_EMU))
        r = p.add_run("\t" + (texto or "").upper())
        _style_run(r, font=FONT_PRIMARY, size=13, color=GREEN_DARK, bold=True)
        _para_border(p, side="bottom", color=GREEN_LIGHT, sz=10, space=8)

    def _setup_footer(self):
        ftr = self.doc.sections[0].footer
        p = ftr.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Twips(TAB_RIGHT), WD_TAB_ALIGNMENT.RIGHT)
        _p_spacing(p, before=40)
        _para_border(p, side="top", color=GREEN_LIGHT, sz=8, space=6)
        r = p.add_run("ComBio Energia\t")
        _style_run(r, font=FONT_SECONDARY, size=8, color=INK_60)
        r = p.add_run("Página ")
        _style_run(r, font=FONT_SECONDARY, size=8, color=INK_60)
        _field(p, "PAGE")
        r = p.add_run(" de ")
        _style_run(r, font=FONT_SECONDARY, size=8, color=INK_60)
        _field(p, "NUMPAGES")

    def _ensure_numbering(self):
        if self._num_id is not None:
            return self._num_id
        from docx.oxml.parser import parse_xml
        try:
            numbering = self.doc.part.numbering_part.element
        except Exception:  # pragma: no cover
            raise RuntimeError("template sem numbering part")
        aid = 900
        numbering.insert(0, parse_xml(BULLET_ABSTRACT_XML.format(aid=aid)))
        nid = 900
        num = _el("w:num", numId=nid)
        num.append(_el("w:abstractNumId", val=aid))
        numbering.append(num)
        self._num_id = nid
        return nid

    # ------------------------------------------------------------ primitivas
    def eyebrow(self, texto):
        p = self.doc.add_paragraph()
        _p_spacing(p, after=140)
        _keep_next(p)
        r = p.add_run(texto.upper())
        _style_run(r, font=FONT_SECONDARY, size=9, color=GREEN_LIGHT, spacing=24)
        return p

    def titulo(self, texto):
        p = self.doc.add_paragraph()
        _p_spacing(p, after=260)
        _keep_next(p)
        r = p.add_run(texto)
        _style_run(r, font=FONT_PRIMARY, size=24, color=GREEN_DARK, bold=True)
        return p

    def orientacao(self, texto):
        """Linha de orientacao/instrucao — Varela Round italico 9,5pt."""
        p = self.doc.add_paragraph()
        _p_spacing(p, after=160)
        r = p.add_run(texto)
        _style_run(r, font=FONT_SECONDARY, size=9.5, color=INK_60, italic=True)
        return p

    def faixa(self, texto):
        """Faixa verde de secao. Em documentos numerados, passe '9. TITULO'."""
        if getattr(self, "_faixa_feita", False):
            self._espaco(220)
        self._faixa_feita = True
        t = self.doc.add_table(rows=1, cols=1)
        self._tbl_base(t, [CONTENT_W], borders=False)
        cell = t.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        _set_shading(tcPr, GREEN_DARK)
        _cell_margins(tcPr, top=90, bottom=90, left=160, right=160)
        _valign(tcPr, "center")
        p = cell.paragraphs[0]
        _p_spacing(p, after=0)
        _keep_next(p)
        r = p.add_run(texto.upper())
        _style_run(r, font=FONT_PRIMARY, size=11, color=WHITE, bold=True, spacing=12)
        self._espaco(120)
        return t

    def subtitulo(self, texto):
        """Subsecao 4.1, 10.2, 12.3 — Kodchasan Bold 12pt verde escuro."""
        p = self.doc.add_paragraph()
        _p_spacing(p, before=120, after=100)
        _keep_next(p)
        r = p.add_run(texto)
        _style_run(r, font=FONT_PRIMARY, size=12, color=GREEN_DARK, bold=True)
        return p

    def paragrafo(self, texto, size=10.5, color=INK_80, bold=False, after=120):
        p = self.doc.add_paragraph()
        _p_spacing(p, after=after)
        self._runs_inline(p, texto, size=size, color=color, bold=bold)
        return p

    def bullets(self, itens, nivel=0):
        nid = self._ensure_numbering()
        for it in itens:
            p = self.doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            numPr.append(_el("w:ilvl", val=nivel))
            numPr.append(_el("w:numId", val=nid))
            _insert_ordered(pPr, numPr)
            p.paragraph_format.left_indent = Twips(360 + nivel * 360)
            _p_spacing(p, after=90)
            self._runs_inline(p, it)
        return None

    def numerado(self, itens):
        """Fluxo numerado (gatilho -> passos -> erros) do escopo tecnico."""
        for i, it in enumerate(itens, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Twips(360)
            p.paragraph_format.first_line_indent = Twips(-360)
            _p_spacing(p, after=90)
            r = p.add_run("%d. " % i)
            _style_run(r, size=10.5, color=GREEN_DARK, bold=True)
            self._runs_inline(p, it)

    def campos(self, linhas, colunas=1):
        """Bloco de campos (rotulo + dica/valor). linhas = [[rotulo, valor], ...]
        colunas=2 monta duas duplas por linha (4 celulas)."""
        ncols = 2 * colunas
        widths = [CONTENT_W // ncols] * ncols
        widths[-1] = CONTENT_W - sum(widths[:-1])
        nrows = (len(linhas) + colunas - 1) // colunas
        t = self.doc.add_table(rows=nrows, cols=ncols)
        self._tbl_base(t, widths)
        for row in t.rows:
            _insert_ordered(row._tr.get_or_add_trPr(), _el("w:cantSplit", val="true"))
        for i in range(nrows):
            fill = WHITE if i % 2 == 0 else INK_05
            for c in range(colunas):
                idx = i * colunas + c
                rotulo, valor = (linhas[idx] + [""])[:2] if idx < len(linhas) else ("", "")
                for j, txt in enumerate((rotulo, valor)):
                    cell = t.cell(i, c * 2 + j)
                    tcPr = cell._tc.get_or_add_tcPr()
                    _set_shading(tcPr, fill)
                    _cell_margins(tcPr, top=130, bottom=130, left=160, right=160)
                    _valign(tcPr, "center")
                    p = cell.paragraphs[0]
                    _p_spacing(p, after=0)
                    if j == 0:
                        r = p.add_run(txt)
                        _style_run(r, size=10.5, color=INK, bold=True)
                    else:
                        self._runs_inline(p, txt, size=10, color=INK_80)
        self._espaco(160)
        return t

    def tabela(self, colunas, linhas, larguras=None, cabecalho=True):
        """Tabela de dados zebrada com cabecalho verde e repeatHeader."""
        n = len(colunas)
        if larguras:
            tot = sum(larguras)
            larguras = [int(w * CONTENT_W / tot) for w in larguras]
            larguras[-1] = CONTENT_W - sum(larguras[:-1])
        else:
            larguras = [CONTENT_W // n] * n
            larguras[-1] = CONTENT_W - sum(larguras[:-1])
        t = self.doc.add_table(rows=len(linhas) + (1 if cabecalho else 0), cols=n)
        self._tbl_base(t, larguras)
        for row in t.rows:
            _insert_ordered(row._tr.get_or_add_trPr(), _el("w:cantSplit", val="true"))
        off = 0
        if cabecalho:
            off = 1
            trPr = t.rows[0]._tr.get_or_add_trPr()
            _insert_ordered(trPr, _el("w:cantSplit", val="true"))
            _insert_ordered(trPr, _el("w:tblHeader", val="true"))
            for j, texto in enumerate(colunas):
                cell = t.cell(0, j)
                tcPr = cell._tc.get_or_add_tcPr()
                _set_shading(tcPr, GREEN_DARK)
                _cell_margins(tcPr, top=110, bottom=110, left=140, right=140)
                _valign(tcPr, "center")
                p = cell.paragraphs[0]
                _p_spacing(p, after=0)
                r = p.add_run(str(texto))
                _style_run(r, size=10, color=WHITE, bold=True)
        for i, linha in enumerate(linhas):
            fill = WHITE if i % 2 == 0 else INK_02
            for j in range(n):
                texto = str(linha[j]) if j < len(linha) and linha[j] is not None else ""
                cell = t.cell(i + off, j)
                tcPr = cell._tc.get_or_add_tcPr()
                _set_shading(tcPr, fill)
                _cell_margins(tcPr, top=110, bottom=110, left=140, right=140)
                _valign(tcPr, "top")
                first = True
                for parte in texto.split("\n"):
                    p = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    _p_spacing(p, after=0)
                    self._runs_inline(p, parte, size=10, color=INK_80)
        self._espaco(160)
        return t

    def callout(self, texto, cor=YELLOW):
        """Callout 'a confirmar com a TI' — shading F2F4F1 + barra esquerda 3pt."""
        t = self.doc.add_table(rows=1, cols=1)
        self._tbl_base(t, [CONTENT_W], borders=False)
        cell = t.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        _set_shading(tcPr, INK_05)
        _cell_margins(tcPr, top=140, bottom=140, left=180, right=180)
        b = OxmlElement("w:tcBorders")
        b.append(_el("w:left", val="single", sz=24, space=0, color=cor))
        tcPr.append(b)
        p = cell.paragraphs[0]
        _p_spacing(p, after=0)
        self._runs_inline(p, texto, size=10, color=INK_80)
        self._espaco(160)
        return t

    def chip_status(self, texto, cor=GREEN_DARK):
        """Chip de status geral do relatorio de entrega."""
        t = self.doc.add_table(rows=1, cols=1)
        self._tbl_base(t, [int(CONTENT_W * 0.55)], borders=False, align_left=True)
        cell = t.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        _set_shading(tcPr, INK_05)
        _cell_margins(tcPr, top=110, bottom=110, left=180, right=180)
        b = OxmlElement("w:tcBorders")
        b.append(_el("w:left", val="single", sz=24, space=0, color=cor))
        tcPr.append(b)
        p = cell.paragraphs[0]
        _p_spacing(p, after=0)
        r = p.add_run("STATUS GERAL: ")
        _style_run(r, size=10, color=INK_60, bold=True, spacing=12)
        r = p.add_run(texto)
        _style_run(r, size=10, color=cor, bold=True)
        self._espaco(160)
        return t

    def codigo(self, texto):
        """Trecho tecnico — Consolas 9,5pt com shading F2F4F1."""
        for linha in texto.split("\n"):
            p = self.doc.add_paragraph()
            _p_spacing(p, after=0)
            _para_shading(p, INK_05)
            p.paragraph_format.left_indent = Twips(120)
            p.paragraph_format.right_indent = Twips(120)
            r = p.add_run(linha or " ")
            _style_run(r, font=FONT_MONO, size=9.5, color=INK)
        self._espaco(160)

    def imagem(self, caminho, largura_px=600, legenda=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _p_spacing(p, after=60)
        p.add_run().add_picture(caminho, width=Emu(int(largura_px * 9525)))
        if legenda:
            c = self.doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _p_spacing(c, after=160)
            r = c.add_run(legenda)
            _style_run(r, font=FONT_SECONDARY, size=8, color=INK_60)

    def regua(self):
        p = self.doc.add_paragraph()
        _p_spacing(p, before=60, after=120)
        _para_border(p, side="bottom", color=GREEN_LIGHT, sz=6, space=1)

    def quebra_pagina(self):
        from docx.enum.text import WD_BREAK
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def indice(self, nomes):
        """Indice das secoes, para documentos longos (3.3 e 3.4).

        Nao usa campo TOC: como as secoes sao faixas (tabelas), e nao paragrafos
        com estilo Heading, um TOC nativo sairia vazio. Aqui o indice e montado
        a partir dos proprios nomes das faixas — sempre correto, sem exigir F9.
        """
        self.faixa("ÍNDICE")
        for nome in nomes:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Twips(160)
            _p_spacing(p, after=60)
            r = p.add_run(nome)
            _style_run(r, size=10.5, color=INK_80)
        self._espaco(200)

    # ------------------------------------------------------------- internos
    def _espaco(self, twips=120):
        """Respiro entre blocos: paragrafo vazio com altura = twips/20 pt."""
        p = self.doc.add_paragraph()
        _p_spacing(p, before=0, after=0, line=1.0)
        r = p.add_run("")
        _style_run(r, size=max(1.0, twips / 20.0), color=WHITE)
        return p

    def _tbl_base(self, table, widths, borders=True, align_left=False):
        table.alignment = WD_TABLE_ALIGNMENT.LEFT if align_left else WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        tblPr = table._tbl.tblPr
        _insert_ordered(tblPr, _el("w:tblW", w=sum(widths), type="dxa"))
        _insert_ordered(tblPr, _el("w:tblLayout", type="fixed"))
        if borders:
            _borders(tblPr, color=INK_20, sz=4,
                     sides=("top", "left", "bottom", "right", "insideH", "insideV"))
        else:
            b = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                b.append(_el("w:" + side, val="nil"))
            _insert_ordered(tblPr, b)
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            table._tbl.remove(grid)
        grid = OxmlElement("w:tblGrid")
        for w in widths:
            grid.append(_el("w:gridCol", w=w))
        tblPr.addnext(grid)
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Twips(widths[i] if i < len(widths) else widths[-1])

    def _runs_inline(self, paragraph, texto, size=10.5, color=INK_80, bold=False):
        """Aplica marcador de status colorido e **negrito** inline."""
        texto = "" if texto is None else str(texto)
        marcador, resto = _split_status(texto)
        if marcador:
            # U+FE0E força apresentação TEXTO (monocromática): sem ele o
            # LibreOffice/Word trocam o ⚠ por emoji colorido, o que o padrão
            # ComBio proíbe.
            r = paragraph.add_run(marcador + "︎")
            _style_run(r, font=FONT_SIMBOLO, size=size,
                       color=STATUS_COLORS[marcador], bold=True)
            r2 = paragraph.add_run(" ")
            _style_run(r2, size=size, color=color)
            texto = resto
        for parte in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
            if not parte:
                continue
            if parte.startswith("**") and parte.endswith("**"):
                r = paragraph.add_run(parte[2:-2])
                _style_run(r, size=size, color=INK, bold=True)
            elif parte.startswith("`") and parte.endswith("`"):
                r = paragraph.add_run(parte[1:-1])
                _style_run(r, font=FONT_MONO, size=size - 1, color=INK)
            else:
                r = paragraph.add_run(parte)
                _style_run(r, size=size, color=color, bold=bold)

    # ------------------------------------------------------------------ save
    def salvar(self, caminho):
        os.makedirs(os.path.dirname(os.path.abspath(caminho)) or ".", exist_ok=True)
        # Carimba a versao da skill nas propriedades do arquivo, para
        # rastrear com qual versao cada documento foi gerado.
        cp = self.doc.core_properties
        cp.category = "ComBio Energia — documento de escopo/entrega"
        cp.comments = "Gerado pela skill %s v%s" % (SKILL_NOME, VERSAO)
        self.doc.save(caminho)
        return caminho
